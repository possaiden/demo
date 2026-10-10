"""Live, formula-driven Excel Gantt.

Every date in the workbook is an Excel formula fed by the yellow cells on the
Inputs sheet, so changing the contract date (or any duration) re-plans the
whole programme and redraws the bars without re-running this script.
"""
from __future__ import annotations

import math
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .schedule import ACTIVITIES, ProjectPlan
from .style import (
    ACTIVITY_COLORS, BAND_FILL, CONTRACT_COLOR, DIAMOND, FONT, GO_LIVE_COLOR, GRID,
    HEADER_FILL, INK, INPUT_FILL, INPUT_FONT, MUTED, STAR, TODAY_COLOR,
)

DATE_FMT = "dd-mmm-yy"
HOLIDAY_ROWS = 200
SPARE_WEEKS = 6


def _font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def _fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


THIN = Side(style="thin", color=GRID)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", indent=1)


def _input_cell(cell, value, number_format=None, note=None):
    cell.value = value
    cell.font = _font(bold=True, color=INPUT_FONT)
    cell.fill = _fill(INPUT_FILL)
    cell.border = BOX
    cell.alignment = CENTER
    if number_format:
        cell.number_format = number_format
    if note:
        cell.comment = Comment(note, "Planner")


# --------------------------------------------------------------------------- Inputs
def _build_inputs(wb: Workbook, plan: ProjectPlan) -> dict:
    ws = wb.create_sheet("Inputs")
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDEFGH", (2, 30, 16, 18, 16, 18, 16, 16)):
        ws.column_dimensions[col].width = width

    ws["B1"] = "Plan Inputs"
    ws["B1"].font = _font(16, bold=True)
    ws["B2"] = ("Edit only the YELLOW cells (blue text). Every date on the 'Plan' sheet is a "
                "formula and updates automatically.")
    ws["B2"].font = _font(9, italic=True, color=MUTED)

    labels = [
        (4, "Project name", plan.project_name, None, None),
        (5, "Client", plan.client_name, None, None),
        (6, "Contract signed date " + STAR, plan.contract_date, DATE_FMT,
         "Source: user input. Default contract date for every process; "
         "override per process in the table below."),
        (7, "Pre-Onboarding (weeks)", plan.pre_onboarding_weeks, "0",
         "Rule from user: Pre-Onboarding = 2 weeks after hiring end date."),
    ]
    for row, label, value, fmt, note in labels:
        ws.cell(row=row, column=2, value=label).font = _font(bold=True)
        _input_cell(ws.cell(row=row, column=3), value, fmt, note)
        if isinstance(value, str):  # long text inputs: widen and left-align
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
            ws.cell(row=row, column=3).alignment = LEFT

    # Vertical rules table
    ws["B9"] = "KT & Guided Production rules by vertical"
    ws["B9"].font = _font(11, bold=True)
    for col, head in zip("BCD", ("Vertical", "KT (weeks)", "Guided Prod. (weeks)")):
        c = ws[f"{col}10"]
        c.value = head
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = _fill(HEADER_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[10].height = 30
    rule_first = 11
    for i, (vertical, rule) in enumerate(plan.vertical_rules.items()):
        r = rule_first + i
        _input_cell(ws.cell(row=r, column=2), vertical)
        _input_cell(ws.cell(row=r, column=3), rule["kt_weeks"], "0")
        _input_cell(ws.cell(row=r, column=4), rule["gp_weeks"], "0")
    rule_last = rule_first + len(plan.vertical_rules) - 1
    ws.cell(row=rule_last + 1, column=2,
            value="Source: user rule — 4 weeks for P2P & OTC, 8 weeks for RTR.").font = \
        _font(8, italic=True, color=MUTED)

    # Process table
    head_row = rule_last + 4
    ws.cell(row=head_row - 1, column=2, value="Processes").font = _font(11, bold=True)
    heads = ("Process", "Vertical", "Contract signed " + STAR, "Hiring (weeks)",
             "Pre-Onboarding (weeks)", "KT (weeks)", "Guided Prod. (weeks)")
    for j, head in enumerate(heads):
        c = ws.cell(row=head_row, column=2 + j, value=head)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = _fill(HEADER_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[head_row].height = 30

    rules_v = f"$B${rule_first}:$B${rule_last}"
    proc_rows = []
    for i, proc in enumerate(plan.processes):
        r = head_row + 1 + i
        proc_rows.append(r)
        _input_cell(ws.cell(row=r, column=2), proc.name)
        _input_cell(ws.cell(row=r, column=3), proc.vertical)
        # Contract date defaults to the project date (formula) — overwrite to override.
        contract_val = ("=ContractDate" if proc.contract_date == plan.contract_date
                        else proc.contract_date)
        _input_cell(ws.cell(row=r, column=4), contract_val, DATE_FMT)
        _input_cell(ws.cell(row=r, column=5), proc.hiring_weeks, "0")
        formulas = (
            "=PreOnboardingWeeks",
            f"=INDEX($C${rule_first}:$C${rule_last},MATCH(C{r},{rules_v},0))",
            f"=INDEX($D${rule_first}:$D${rule_last},MATCH(C{r},{rules_v},0))",
        )
        for j, f in enumerate(formulas):
            c = ws.cell(row=r, column=6 + j, value=f)
            c.font = _font()
            c.border = BOX
            c.alignment = CENTER
            c.number_format = "0"

    first, last = proc_rows[0], proc_rows[-1]
    dv = DataValidation(type="whole", operator="between", formula1="8", formula2="10",
                        errorStyle="warning", showErrorMessage=True,
                        errorTitle="Hiring duration",
                        error="Standard hiring timeline is 8–10 weeks. Keep this value?")
    ws.add_data_validation(dv)
    dv.add(f"E{first}:E{last}")
    verticals = ",".join(plan.vertical_rules)
    dv2 = DataValidation(type="list", formula1=f'"{verticals}"', allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add(f"C{first}:C{last}")

    notes_row = last + 2
    notes = [
        "How dates are calculated (working days Mon–Fri, holidays on the 'Holidays' sheet are skipped):",
        f"  {STAR} Contract Signed — from the table above (defaults to the project contract date).",
        "  Hiring — starts next working day after contract; 8–10 weeks standard.",
        "  Pre-Onboarding — starts next working day after hiring ends; 2 weeks.",
        "  Knowledge Transfer — after Pre-Onboarding; 4 weeks P2P/OTC, 8 weeks RTR.",
        "  Guided Production — after KT; 4 weeks P2P/OTC, 8 weeks RTR.",
        f"  {DIAMOND} Go-Live — next working day after Guided Production ends.",
        "  1 week = 5 working days. Pre-Onboarding / KT / Guided Prod. columns in the process "
        "table are formulas — change the rules above instead.",
    ]
    for k, text in enumerate(notes):
        ws.cell(row=notes_row + k, column=2, value=text).font = \
            _font(9, bold=(k == 0), color=INK if k == 0 else MUTED)

    wb.defined_names["ContractDate"] = DefinedName("ContractDate", attr_text="Inputs!$C$6")
    wb.defined_names["PreOnboardingWeeks"] = DefinedName(
        "PreOnboardingWeeks", attr_text="Inputs!$C$7")
    return {"proc_rows": proc_rows}


def _build_holidays(wb: Workbook, plan: ProjectPlan):
    ws = wb.create_sheet("Holidays")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 34
    ws["A1"] = "Holiday date"
    ws["B1"] = "Description (optional)"
    for c in (ws["A1"], ws["B1"]):
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = _fill(HEADER_FILL)
    ws["D1"] = "Add public / client holidays in column A. They are skipped by every plan date."
    ws["D1"].font = _font(9, italic=True, color=MUTED)
    for i in range(HOLIDAY_ROWS):
        c = ws.cell(row=2 + i, column=1)
        c.number_format = DATE_FMT
        c.fill = _fill(INPUT_FILL)
        c.font = _font(color=INPUT_FONT)
        if i < len(plan.holidays):
            c.value = plan.holidays[i]
    wb.defined_names["Holidays"] = DefinedName(
        "Holidays", attr_text=f"Holidays!$A$2:$A${HOLIDAY_ROWS + 1}")


# --------------------------------------------------------------------------- Plan
def _build_plan_sheet(wb: Workbook, plan: ProjectPlan, inputs: dict):
    ws = wb.create_sheet("Plan", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90

    first_week = plan.start - timedelta(days=plan.start.weekday())
    n_weeks = math.ceil((plan.final_go_live - first_week).days / 7) + 1 + SPARE_WEEKS
    n_weeks = max(n_weeks, 21)
    TL0 = 7  # column G = first timeline week
    tl_cols = [get_column_letter(TL0 + i) for i in range(n_weeks)]
    tl_first, tl_last = tl_cols[0], tl_cols[-1]

    widths = {"A": 1.5, "B": 30, "C": 7, "D": 11.5, "E": 11.5, "F": 1.5}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for col in tl_cols:
        ws.column_dimensions[col].width = 4.6

    proc_rows = inputs["proc_rows"]

    # --- Title ---------------------------------------------------------------
    ws["B1"] = '=Inputs!C4&" — Transition Plan"'
    ws["B1"].font = _font(18, bold=True)
    ws["B2"] = '="Client: "&Inputs!C5'
    ws["B2"].font = _font(10, color=MUTED)
    ws.row_dimensions[1].height = 28

    # --- Go-Live KPI tiles (rows 4–6) -----------------------------------------
    # Rows of each process block are computed below; fill tiles after layout.
    HEAD_MONTH, HEAD_WEEK = 8, 9
    row = HEAD_WEEK + 1
    blocks = []
    for i, proc in enumerate(plan.processes):
        r_head = row
        r_contract = row + 1
        r_phases = [row + 2 + k for k in range(len(ACTIVITIES))]
        r_golive = row + 2 + len(ACTIVITIES)
        blocks.append((i, proc, r_head, r_contract, r_phases, r_golive))
        row = r_golive + 2
    last_row = row - 2

    def tile(col_start, span, title, value_formula, sub_formula, accent):
        c0 = get_column_letter(col_start)
        c1 = get_column_letter(col_start + span - 1)
        for r, val, fnt, fmt in (
            (4, title, _font(9, bold=True, color=MUTED), None),
            (5, value_formula, _font(16, bold=True, color=accent), DATE_FMT),
            (6, sub_formula, _font(8, color=MUTED), None),
        ):
            ws.merge_cells(f"{c0}{r}:{c1}{r}")
            cell = ws[f"{c0}{r}"]
            cell.value = val
            cell.font = fnt
            cell.alignment = CENTER
            if fmt:
                cell.number_format = fmt
            for cc in range(col_start, col_start + span):
                ws.cell(row=r, column=cc).fill = _fill(BAND_FILL)
    ws.row_dimensions[5].height = 26

    golive_cells = [f"$D${b[5]}" for b in blocks]
    contract_cells = [f"$D${b[3]}" for b in blocks]
    # Final go-live tile over columns B..E region (merged B:E)
    ws.merge_cells("B4:E4"); ws.merge_cells("B5:E5"); ws.merge_cells("B6:E6")
    ws["B4"] = DIAMOND + "  ALL PROCESSES LIVE BY"
    ws["B4"].font = _font(9, bold=True, color="FFFFFF")
    ws["B5"] = f"=MAX({','.join(golive_cells)})"
    ws["B5"].font = _font(18, bold=True, color="FFFFFF")
    ws["B5"].number_format = "dd-mmm-yyyy"
    ws["B6"] = (f'=ROUND((B5-MIN({",".join(contract_cells)}))/7,0)'
                f'&" weeks from contract signature"')
    ws["B6"].font = _font(8, color="D1D5DB")
    for r in (4, 5, 6):
        ws[f"B{r}"].alignment = CENTER
        for col in "BCDE":
            ws[f"{col}{r}"].fill = _fill(HEADER_FILL)

    span = 6
    for i, proc, r_head, r_contract, r_phases, r_golive in blocks:
        tile(TL0 + i * (span + 1), span,
             f"=UPPER(Inputs!C{proc_rows[i]})&\" {DIAMOND} GO-LIVE\"",
             f"=D{r_golive}",
             f'=ROUND((D{r_golive}-D{r_contract})/7,0)&" wks from contract"',
             GO_LIVE_COLOR)

    # --- Timeline headers ------------------------------------------------------
    for col, text in zip("BCDE", ("Process / Activity", "Weeks", "Start", "End")):
        c = ws[f"{col}{HEAD_WEEK}"]
        c.value = text
        c.font = _font(10, bold=True, color="FFFFFF")
        c.fill = _fill(HEADER_FILL)
        c.alignment = CENTER if col != "B" else LEFT
    ws[f"B{HEAD_MONTH}"] = "Timeline: week commencing (Mon)"
    ws[f"B{HEAD_MONTH}"].font = _font(8, italic=True, color=MUTED)

    all_contracts = ",".join(f"Inputs!$D${r}" for r in proc_rows)
    for i, col in enumerate(tl_cols):
        wk = ws[f"{col}{HEAD_WEEK}"]
        mo = ws[f"{col}{HEAD_MONTH}"]
        if i == 0:
            wk.value = f"=MIN({all_contracts})-WEEKDAY(MIN({all_contracts}),3)"
            mo.value = f"={col}{HEAD_WEEK}"
        else:
            prev = tl_cols[i - 1]
            wk.value = f"={prev}{HEAD_WEEK}+7"
            mo.value = f'=IF(MONTH({col}{HEAD_WEEK})<>MONTH({prev}{HEAD_WEEK}),{col}{HEAD_WEEK},"")'
        wk.number_format = "dd"
        wk.font = _font(8, bold=True, color="FFFFFF")
        wk.fill = _fill(HEADER_FILL)
        wk.alignment = CENTER
        mo.number_format = "mmm"
        mo.font = _font(9, bold=True, color=INK)
        mo.alignment = Alignment(horizontal="left", vertical="bottom")

    # --- Process blocks ----------------------------------------------------------
    star_font = _font(13, bold=True, color=CONTRACT_COLOR)
    diamond_font = _font(12, bold=True, color=GO_LIVE_COLOR)
    for i, proc, r_head, r_contract, r_phases, r_golive in blocks:
        ir = proc_rows[i]
        # Header band
        ws[f"B{r_head}"] = f'=Inputs!B{ir}&"  ("&Inputs!C{ir}&")"'
        ws[f"B{r_head}"].font = _font(11, bold=True)
        ws[f"C{r_head}"] = f"=ROUND((E{r_head}-D{r_head})/7,0)"
        ws[f"D{r_head}"] = f"=D{r_contract}"
        ws[f"E{r_head}"] = f"=D{r_golive}"
        for col in ["B", "C", "D", "E"] + tl_cols:
            ws[f"{col}{r_head}"].fill = _fill(BAND_FILL)
        for col in "CDE":
            ws[f"{col}{r_head}"].font = _font(9, bold=True, color=MUTED)
        ws.row_dimensions[r_head].height = 20

        # Contract milestone
        ws[f"B{r_contract}"] = f"{STAR}  Contract Signed"
        ws[f"B{r_contract}"].font = _font(10, bold=True, color=INK)
        ws[f"D{r_contract}"] = f"=Inputs!D{ir}"
        ws[f"E{r_contract}"] = f"=D{r_contract}"

        # Phases
        duration_src = {"hiring": "E", "pre_onboarding": "F", "kt": "G", "gp": "H"}
        prev_end = f"E{r_contract}"
        for (key, label), r in zip(ACTIVITIES, r_phases):
            ws[f"B{r}"] = f"      {label}"
            ws[f"B{r}"].font = _font(10)
            ws[f"C{r}"] = f"=Inputs!{duration_src[key]}{ir}"
            ws[f"D{r}"] = f"=WORKDAY({prev_end},1,Holidays)"
            ws[f"E{r}"] = f"=WORKDAY(D{r},C{r}*5-1,Holidays)"
            prev_end = f"E{r}"
            ws.conditional_formatting.add(
                f"{tl_first}{r}:{tl_last}{r}",
                FormulaRule(formula=[f"AND({tl_first}${HEAD_WEEK}<=$E{r},"
                                     f"{tl_first}${HEAD_WEEK}+6>=$D{r})"],
                            fill=_fill(ACTIVITY_COLORS[key])))

        # Go-live milestone
        ws[f"B{r_golive}"] = f"{DIAMOND}  Go-Live"
        ws[f"B{r_golive}"].font = _font(10, bold=True, color=GO_LIVE_COLOR)
        ws[f"D{r_golive}"] = f"=WORKDAY({prev_end},1,Holidays)"
        ws[f"E{r_golive}"] = f"=D{r_golive}"

        for r, symbol, fnt in ((r_contract, STAR, star_font), (r_golive, DIAMOND, diamond_font)):
            for col in tl_cols:
                c = ws[f"{col}{r}"]
                c.value = (f'=IF(AND($D{r}>={col}${HEAD_WEEK},$D{r}<{col}${HEAD_WEEK}+7),'
                           f'"{symbol}","")')
                c.font = fnt
                c.alignment = CENTER

        for r in [r_contract, *r_phases, r_golive]:
            ws[f"C{r}"].number_format = "0"
            for col in "CDE":
                ws[f"{col}{r}"].alignment = CENTER
                if not ws[f"{col}{r}"].font.bold:
                    ws[f"{col}{r}"].font = _font(10)
            ws.row_dimensions[r].height = 18
        for col in "DE":
            for r in range(r_head, r_golive + 1):
                ws[f"{col}{r}"].number_format = DATE_FMT
        ws[f"C{r_contract}"] = None
        ws[f"C{r_golive}"] = None
        for r in range(r_head, r_golive + 1):
            for col in tl_cols:
                ws[f"{col}{r}"].border = Border(left=Side(style="hair", color=GRID),
                                                bottom=Side(style="hair", color=GRID))

    # --- Today marker ------------------------------------------------------------
    today_side = Side(style="medium", color=TODAY_COLOR)
    ws.conditional_formatting.add(
        f"{tl_first}{HEAD_WEEK + 1}:{tl_last}{last_row}",
        FormulaRule(formula=[f"AND(TODAY()>={tl_first}${HEAD_WEEK},"
                             f"TODAY()<{tl_first}${HEAD_WEEK}+7)"],
                    border=Border(left=today_side, right=today_side)))
    ws.conditional_formatting.add(
        f"{tl_first}{HEAD_WEEK}:{tl_last}{HEAD_WEEK}",
        FormulaRule(formula=[f"AND(TODAY()>={tl_first}${HEAD_WEEK},"
                             f"TODAY()<{tl_first}${HEAD_WEEK}+7)"],
                    fill=_fill(TODAY_COLOR)))

    # --- Legend --------------------------------------------------------------------
    leg = last_row + 2
    ws[f"B{leg}"] = "Legend"
    ws[f"B{leg}"].font = _font(9, bold=True, color=MUTED)
    items = [(STAR, "Contract Signed", None, CONTRACT_COLOR)]
    items += [(None, label, ACTIVITY_COLORS[key], None) for key, label in ACTIVITIES]
    items += [(DIAMOND, "Go-Live", None, GO_LIVE_COLOR), (None, "Today", TODAY_COLOR, None)]
    leg_col = TL0
    for symbol, label, fill, fcolor in items:
        swatch = ws.cell(row=leg, column=leg_col)
        if fill:
            swatch.fill = _fill(fill)
        if symbol:
            swatch.value = symbol
            swatch.font = _font(12, bold=True, color=fcolor)
            swatch.alignment = CENTER
        ws.cell(row=leg, column=leg_col + 1, value=label).font = _font(9, color=INK)
        leg_col += 2 + math.ceil(len(label) / 4.5)
    ws[f"B{leg + 1}"] = ("Working days (Mon–Fri); 1 week = 5 working days. Edit dates & durations "
                         "on the 'Inputs' sheet — this chart redraws automatically.")
    ws[f"B{leg + 1}"].font = _font(8, italic=True, color=MUTED)

    ws.freeze_panes = f"{tl_first}{HEAD_WEEK + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{tl_last}{leg + 1}"
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4


def export_excel(plan: ProjectPlan, path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    inputs = _build_inputs(wb, plan)
    _build_holidays(wb, plan)
    _build_plan_sheet(wb, plan, inputs)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # Excel computes every date on open
    wb.save(path)
