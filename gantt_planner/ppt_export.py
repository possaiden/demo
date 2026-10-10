"""Editable PowerPoint: executive summary, Gantt, key-dates table.

Everything is a native shape/table (no images), so bars can be dragged,
recoloured or relabelled directly in PowerPoint.
"""
from __future__ import annotations

from datetime import date, timedelta

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .schedule import ACTIVITIES, ProjectPlan
from .style import (
    ACTIVITY_COLORS, BAND_FILL, CONTRACT_COLOR, DIAMOND, GO_LIVE_COLOR, GRID, INK,
    MUTED, STAR, TODAY_COLOR,
)

FONT = "Calibri"
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5


def _rgb(hex_):
    return RGBColor.from_string(hex_)


def _d(d: date, fmt="%d-%b-%y") -> str:
    return d.strftime(fmt)


def _text(slide, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, italic=False, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    runs = text if isinstance(text, list) else [(text, {})]
    p = tf.paragraphs[0]
    p.alignment = align
    for t, opts in runs:
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = FONT
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.italic = opts.get("italic", italic)
        f.color.rgb = _rgb(opts.get("color", color))
    if name:
        box.name = name
    return box


def _rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(fill)
    if line:
        s.line.color.rgb = _rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    _strip_style(s)
    if name:
        s.name = name
    return s


def _strip_style(shape):
    """Drop the theme style reference (it adds a soft shadow); fill/line are explicit."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _shape_text(shape, text, size=10, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = _rgb(color)


def _line(slide, x1, y1, x2, y2, color=GRID, width=0.75, dash=None, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = _rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    _strip_style(c)
    if name:
        c.name = name
    return c


def _title(slide, title, subtitle=None):
    _text(slide, MARGIN, 0.35, SLIDE_W - 2 * MARGIN, 0.6, title, size=28, bold=True,
          name="Title")
    if subtitle:
        _text(slide, MARGIN, 0.95, SLIDE_W - 2 * MARGIN, 0.35, subtitle, size=14,
              color=MUTED, name="Subtitle")


def _footer(slide, plan: ProjectPlan, page: int):
    _text(slide, MARGIN, SLIDE_H - 0.35, 9, 0.25,
          f"{plan.project_name}  |  {plan.client_name}  |  Working days Mon–Fri; "
          f"1 week = 5 working days", size=10, color=MUTED)
    _text(slide, SLIDE_W - MARGIN - 1, SLIDE_H - 0.35, 1, 0.25, str(page), size=10,
          color=MUTED, align=PP_ALIGN.RIGHT)


# --------------------------------------------------------------------------- slide 1
def _summary_slide(prs, plan: ProjectPlan):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    final = plan.final_go_live
    weeks_total = round((final - plan.start).days / 7)
    _title(s, f"All processes live by {final:%d %b %Y} — {weeks_total} weeks from contract",
           f"{STAR} Contract signed {plan.contract_date:%d %b %Y}   ·   "
           f"{len(plan.processes)} processes   ·   phased go-live by vertical")

    n = len(plan.processes)
    gap = 0.3
    card_w = (SLIDE_W - 2 * MARGIN - gap * (n - 1)) / n
    card_y, card_h = 1.65, 4.25
    for i, p in enumerate(plan.processes):
        x = MARGIN + i * (card_w + gap)
        _rect(s, x, card_y, card_w, card_h, BAND_FILL, MSO_SHAPE.ROUNDED_RECTANGLE,
              name=f"Card {p.vertical}").adjustments[0] = 0.05
        ix, iw = x + 0.3, card_w - 0.6
        _text(s, ix, card_y + 0.25, iw, 0.45, p.name, size=22, bold=True)
        tag = _rect(s, ix, card_y + 0.75, 0.75, 0.3, INK, MSO_SHAPE.ROUNDED_RECTANGLE)
        _shape_text(tag, p.vertical, size=11)
        _text(s, ix, card_y + 1.3, iw, 0.3, f"{DIAMOND}  GO-LIVE", size=12, bold=True,
              color=GO_LIVE_COLOR)
        _text(s, ix, card_y + 1.6, iw, 0.8, f"{p.go_live:%d %b %Y}", size=40, bold=True,
              color=GO_LIVE_COLOR)
        _text(s, ix, card_y + 2.4, iw, 0.3,
              f"{p.go_live:%A}  ·  {round(p.weeks_to_go_live)} weeks from contract",
              size=13, color=MUTED)

        # Proportional phase strip
        strip_y = card_y + 3.0
        total = sum(ph.weeks for ph in p.phases)
        cx = ix
        for ph in p.phases:
            w = iw * ph.weeks / total
            seg = _rect(s, cx, strip_y, w, 0.32, ACTIVITY_COLORS[ph.key],
                        name=f"{p.vertical} {ph.label}")
            _shape_text(seg, f"{ph.weeks:g}w", size=10,
                        color=INK if ph.key in ("pre_onboarding", "kt") else "FFFFFF")
            cx += w
        _text(s, ix, strip_y + 0.45, iw, 0.6,
              "Hiring → Pre-Onboarding → KT → Guided Production",
              size=11, color=MUTED)

    crit = max(plan.processes, key=lambda p: p.go_live)
    kt = next(ph for ph in crit.phases if ph.key == "kt")
    gp = next(ph for ph in crit.phases if ph.key == "gp")
    _text(s, MARGIN, 6.2, SLIDE_W - 2 * MARGIN, 0.6, [
        ("Critical path: ", {"bold": True}),
        (f"{crit.name} ({crit.vertical}) — {crit.hiring_weeks:g}-week hiring plus "
         f"{kt.weeks:g}-week KT and {gp.weeks:g}-week Guided Production. "
         f"Any slip in contract signature moves every go-live date one-for-one.", {}),
    ], size=14)
    _footer(s, plan, 1)


# --------------------------------------------------------------------------- slide 2
def _gantt_slide(prs, plan: ProjectPlan):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, "Timeline: contract to go-live by process",
           f"{STAR} Contract signed  ·  bars = working weeks per phase  ·  "
           f"{DIAMOND} Go-Live = next working day after Guided Production")

    t0 = plan.start - timedelta(days=plan.start.weekday())
    end = plan.final_go_live + timedelta(days=14)
    t1 = date(end.year + (end.month == 12), end.month % 12 + 1, 1)  # month after end
    span = (t1 - t0).days
    label_w = 2.6
    cx0 = MARGIN + label_w + 0.1
    cx1 = SLIDE_W - MARGIN
    cw = cx1 - cx0

    def X(d: date) -> float:
        return cx0 + cw * (d - t0).days / span

    head_y = 1.45
    head_rows = 0.38
    body_y = 1.95
    row_h, proc_h, gap = 0.25, 0.36, 0.1
    body_h = len(plan.processes) * (proc_h + len(ACTIVITIES) * row_h + gap) - gap
    body_bottom = body_y + body_h

    # Month header + gridlines
    m = date(t0.year, t0.month, 1)
    while m < t1:
        nxt = date(m.year + (m.month == 12), m.month % 12 + 1, 1)
        a, b = max(m, t0), min(nxt, t1)
        if (b - a).days >= 10:
            _text(s, X(a), head_y, X(b) - X(a), head_rows,
                  [(f"{a:%b}", {"bold": True}), (f" {a:%y}", {"color": MUTED, "size": 10})],
                  size=12, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if m > t0:
            _line(s, X(m), head_y, X(m), body_bottom, color=GRID)
        m = nxt
    _line(s, cx0, head_y + head_rows + 0.04, cx1, head_y + head_rows + 0.04,
          color=MUTED, width=1)

    y = body_y
    for p in plan.processes:
        _rect(s, MARGIN, y, SLIDE_W - 2 * MARGIN, proc_h, BAND_FILL,
                     name=f"{p.vertical} band")
        _text(s, MARGIN + 0.1, y, label_w, proc_h,
              [(p.name, {"bold": True}), (f"  {p.vertical}", {"color": MUTED, "size": 11})],
              size=13, anchor=MSO_ANCHOR.MIDDLE)
        # Contract star
        sz = 0.26
        _rect(s, X(p.contract_date) - sz / 2, y + (proc_h - sz) / 2, sz, sz,
                     CONTRACT_COLOR, MSO_SHAPE.STAR_5_POINT, name=f"{p.vertical} contract star")
        # Go-live diamond + label
        dsz = 0.24
        _rect(s, X(p.go_live) - dsz / 2, y + (proc_h - dsz) / 2, dsz, dsz, GO_LIVE_COLOR,
              MSO_SHAPE.DIAMOND, name=f"{p.vertical} go-live")
        label_w_in = 1.5
        right_room = cx1 - (X(p.go_live) + dsz / 2 + 0.06)
        if right_room >= label_w_in:
            lx_, align = X(p.go_live) + dsz / 2 + 0.06, PP_ALIGN.LEFT
        else:
            lx_, align = X(p.go_live) - dsz / 2 - 0.06 - label_w_in, PP_ALIGN.RIGHT
        _text(s, lx_, y, label_w_in, proc_h, f"Go-Live {p.go_live:%d %b %y}", size=11,
              bold=True, color=GO_LIVE_COLOR, anchor=MSO_ANCHOR.MIDDLE, align=align)
        y += proc_h
        for ph in p.phases:
            _text(s, MARGIN + 0.3, y, label_w - 0.2, row_h, ph.label, size=11,
                  color=INK, anchor=MSO_ANCHOR.MIDDLE)
            bx = X(ph.start)
            bw = max(X(ph.end + timedelta(days=1)) - bx, 0.05)
            bar = _rect(s, bx, y + 0.04, bw, row_h - 0.08, ACTIVITY_COLORS[ph.key],
                        MSO_SHAPE.ROUNDED_RECTANGLE, name=f"{p.vertical} {ph.label} bar")
            bar.adjustments[0] = 0.25
            if bw > 0.4:
                _shape_text(bar, f"{ph.weeks:g} wks", size=9,
                            color=INK if ph.key in ("pre_onboarding", "kt") else "FFFFFF")
            y += row_h
        y += gap

    today = date.today()
    if t0 <= today <= t1:
        _line(s, X(today), head_y + head_rows, X(today), body_bottom, color=TODAY_COLOR,
              width=1.5, dash=MSO_LINE_DASH_STYLE.DASH, name="Today")
        _text(s, X(today) - 0.4, body_bottom + 0.02, 0.8, 0.2, "Today", size=10,
              bold=True, color=TODAY_COLOR, align=PP_ALIGN.CENTER)

    # Legend
    lx, ly = MARGIN, SLIDE_H - 0.8
    items = [("star", "Contract Signed", CONTRACT_COLOR)]
    items += [("bar", label, ACTIVITY_COLORS[key]) for key, label in ACTIVITIES]
    items += [("diamond", "Go-Live", GO_LIVE_COLOR)]
    for kind, label, color in items:
        shape = {"star": MSO_SHAPE.STAR_5_POINT, "diamond": MSO_SHAPE.DIAMOND,
                 "bar": MSO_SHAPE.ROUNDED_RECTANGLE}[kind]
        w = 0.4 if kind == "bar" else 0.2
        _rect(s, lx, ly + 0.03, w, 0.2, color, shape)
        lx += w + 0.08
        tw = 0.075 * len(label) + 0.1
        _text(s, lx, ly, tw, 0.26, label, size=11, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        lx += tw + 0.25
    _footer(s, plan, 2)


# --------------------------------------------------------------------------- slide 3
def _dates_slide(prs, plan: ProjectPlan):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, "Key dates and planning rules",
           "All dates are working days; edit the Excel 'Inputs' sheet to re-plan")
    heads = ["Process", f"{STAR} Contract", "Hiring", "Pre-Onboarding",
             "Knowledge Transfer", "Guided Production", f"{DIAMOND} Go-Live"]
    rows, cols = len(plan.processes) + 1, len(heads)
    widths = [2.4, 1.3, 1.75, 1.75, 1.75, 1.75, 1.633]
    tbl = s.shapes.add_table(rows, cols, Inches(MARGIN), Inches(1.6),
                             Inches(sum(widths)), Inches(0.5 * rows)).table
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)

    def cell(r, c, text, bold=False, color=INK, fill=None, size=13):
        ce = tbl.cell(r, c)
        ce.text = ""
        ce.margin_left = ce.margin_right = Inches(0.08)
        ce.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = ce.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)
        ce.fill.solid()
        ce.fill.fore_color.rgb = _rgb(fill or "FFFFFF")

    for j, h in enumerate(heads):
        cell(0, j, h, bold=True, color="FFFFFF", fill=INK, size=12)
    for i, p in enumerate(plan.processes, start=1):
        fill = BAND_FILL if i % 2 == 0 else "FFFFFF"
        cell(i, 0, f"{p.name} ({p.vertical})", bold=True, fill=fill)
        cell(i, 1, _d(p.contract_date), fill=fill)
        for j, ph in enumerate(p.phases, start=2):
            cell(i, j, f"{ph.start:%d-%b} → {ph.end:%d-%b}  ({ph.weeks:g}w)", fill=fill,
                 size=11)
        cell(i, 6, _d(p.go_live), bold=True, color=GO_LIVE_COLOR, fill=fill)

    rules_y = 1.6 + 0.5 * rows + 0.5
    _text(s, MARGIN, rules_y, 6, 0.35, "Planning rules", size=16, bold=True)
    rule_lines = [
        f"{STAR} Contract Signed — project date (can be overridden per process)",
        "Hiring — 8 to 10 weeks standard, from the next working day after contract",
        f"Pre-Onboarding — {plan.pre_onboarding_weeks:g} weeks after hiring end date",
    ]
    kt_txt = ", ".join(f"{v} {r['kt_weeks']}w" for v, r in plan.vertical_rules.items())
    gp_txt = ", ".join(f"{v} {r['gp_weeks']}w" for v, r in plan.vertical_rules.items())
    rule_lines += [
        f"Knowledge Transfer — {kt_txt}",
        f"Guided Production — {gp_txt}",
        f"{DIAMOND} Go-Live — next working day after Guided Production ends",
    ]
    box = s.shapes.add_textbox(Inches(MARGIN), Inches(rules_y + 0.45), Inches(8.5),
                               Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for k, line in enumerate(rule_lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(14)
        r.font.color.rgb = _rgb(INK)

    # Swatch key for colours on the right
    kx, ky = 9.6, rules_y + 0.5
    _text(s, kx, rules_y, 3, 0.35, "Colour key", size=16, bold=True)
    for k, (key, label) in enumerate(ACTIVITIES):
        _rect(s, kx, ky + k * 0.42, 0.45, 0.24, ACTIVITY_COLORS[key],
              MSO_SHAPE.ROUNDED_RECTANGLE)
        _text(s, kx + 0.6, ky + k * 0.42 - 0.03, 2.5, 0.3, label, size=14,
              anchor=MSO_ANCHOR.MIDDLE)
    if plan.holidays:
        _text(s, MARGIN, SLIDE_H - 0.75, 12, 0.3,
              "Holidays excluded: " + ", ".join(_d(h, "%d-%b-%y") for h in plan.holidays),
              size=10, color=MUTED, italic=True)
    _footer(s, plan, 3)


def export_ppt(plan: ProjectPlan, path) -> None:
    prs = Presentation()
    prs.slide_width = Emu(int(SLIDE_W * 914400))
    prs.slide_height = Emu(int(SLIDE_H * 914400))
    prs.core_properties.title = f"{plan.project_name} — Project Plan"
    _summary_slide(prs, plan)
    _gantt_slide(prs, plan)
    _dates_slide(prs, plan)
    prs.save(path)
